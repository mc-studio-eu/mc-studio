from __future__ import annotations

import csv
import re
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path("outputs/019fe5d0-175f-7c11-ada2-df5e2754bb6a/validation-mails-brevo-startups-lyon.xlsx")
OUT_DIR = Path("outputs/rentree-cold-email-2026/brevo")
CONTACTS_OUT = OUT_DIR / "contacts-startups-lyon.csv"
PREVIEW_OUT = OUT_DIR / "contacts-startups-lyon-preview.csv"


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def mask_email(email: str) -> str:
    if "@" not in email:
        return email
    local, domain = email.split("@", 1)
    return f"{local[:1]}***@{domain}"


def main() -> None:
    workbook_path = SOURCE
    if not workbook_path.exists():
        raise SystemExit(f"Missing source workbook: {workbook_path}")

    wb = load_workbook(workbook_path, read_only=False, data_only=True)
    ws = wb["Mails réécrits"]
    headers = [clean(cell.value) for cell in ws[1]]
    index = {header: position for position, header in enumerate(headers)}

    required = [
        "Jour",
        "Date",
        "Priorité",
        "Startup",
        "Email",
        "Salutation",
        "Objet",
        "Petite introduction",
        "Observation précise",
        "Idée de landing",
        "Offre / CTA",
        "Site source",
    ]
    missing = [name for name in required if name not in index]
    if missing:
        raise SystemExit(f"Missing columns: {', '.join(missing)}")

    rows: list[dict[str, str]] = []
    seen: set[str] = set()
    email_re = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

    for source_row in ws.iter_rows(min_row=2, values_only=True):
        email = clean(source_row[index["Email"]]).lower()
        if not email or not email_re.match(email) or email in seen:
            continue
        seen.add(email)

        rows.append(
            {
                "EMAIL": email,
                "FIRSTNAME": "",
                "LASTNAME": "",
                "STARTUP": clean(source_row[index["Startup"]]),
                "SALUTATION": clean(source_row[index["Salutation"]]),
                "SUBJECT_LINE": clean(source_row[index["Objet"]]),
                "INTRO": clean(source_row[index["Petite introduction"]]),
                "OBSERVATION": clean(source_row[index["Observation précise"]]),
                "LANDING_IDEA": clean(source_row[index["Idée de landing"]]),
                "OFFER_CTA": "Je realise chez MC Studio des landing pages et sites complets pour clarifier l'offre, renforcer la confiance et generer plus de demandes qualifiees. Une landing page demarre a 1 890 EUR HT et un site 5 pages a 2 990 EUR HT.",
                "SOURCE_SITE": clean(source_row[index["Site source"]]),
                "PRIORITY": clean(source_row[index["Priorité"]]),
                "SEQUENCE_DAY": clean(source_row[index["Jour"]]),
                "DATA_ORIGIN": "Base existante MC Studio - validation-mails-brevo-startups-lyon.xlsx",
                "STATUS": "ready_for_brevo_review",
            }
        )

    fields = list(rows[0].keys()) if rows else []
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    with CONTACTS_OUT.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    preview_rows = [{**row, "EMAIL": mask_email(row["EMAIL"])} for row in rows]
    with PREVIEW_OUT.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=fields)
        writer.writeheader()
        writer.writerows(preview_rows)

    print(f"Extracted {len(rows)} unique contacts")
    print(CONTACTS_OUT)
    print(PREVIEW_OUT)


if __name__ == "__main__":
    main()
